from __future__ import annotations

import io
import json
from datetime import datetime
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.infrastructure.models import RunModel, RunReportItemModel, RunReportModel


RUN_STATUS_LABELS = {
    "QUEUED": "排队中",
    "STARTING": "正在启动",
    "RUNNING": "运行中",
    "STOPPING": "停止中",
    "SUCCEEDED": "已完成",
    "FAILED": "执行失败",
    "STOPPED": "已停止",
    "TIMED_OUT": "已超时",
    "INTERRUPTED": "已中断",
}
REPORT_STATUS_LABELS = {"OPEN": "生成中", "COMPLETE": "完整", "INCOMPLETE": "未完整"}
ITEM_STATUS_LABELS = {
    "SUCCESS": "成功",
    "FAILED": "失败",
    "NO_DATA": "无数据",
    "SKIPPED": "已跳过",
    "UNFINISHED": "未完成",
}


def build_report_workbook(
    *,
    run: RunModel,
    customer_name: str,
    report: RunReportModel,
    schema: dict[str, Any],
    items: list[RunReportItemModel],
    timezone_name: str,
) -> bytes:
    timezone = ZoneInfo(timezone_name)
    unreported = max((report.expected_total or report.reported_total) - report.reported_total, 0)
    total = max(report.expected_total or 0, report.reported_total)
    unfinished = report.unfinished_count + unreported

    workbook = Workbook()
    summary = workbook.active
    summary.title = "执行汇总"
    summary.sheet_view.showGridLines = False
    summary.freeze_panes = "A2"
    summary.append(["项目", "内容"])
    summary_rows = [
        ("客户", customer_name),
        ("功能", run.feature_name),
        ("请求 ID", run.request_id),
        ("代码版本", run.version_number),
        ("触发方式", "手动运行" if run.trigger_source == "MANUAL" else "定时运行"),
        ("平台执行状态", RUN_STATUS_LABELS.get(run.status, run.status)),
        ("报表完整性", REPORT_STATUS_LABELS.get(report.status, report.status)),
        ("数据源文件", run.data_source_filename or "不使用数据源"),
        ("进入队列", _format_timestamp(run.queued_at, timezone)),
        ("开始时间", _format_timestamp(run.started_at, timezone)),
        ("完成时间", _format_timestamp(run.finished_at, timezone)),
        ("总数", total),
        ("成功", report.success_count),
        ("失败", report.failed_count),
        ("无数据", report.no_data_count),
        ("已跳过", report.skipped_count),
        ("未完成", unfinished),
        ("已形成明细", report.reported_total),
    ]
    for key, value in summary_rows:
        summary.append([key, _safe_cell(value)])
    _style_header(summary, 2)
    summary.column_dimensions["A"].width = 19
    summary.column_dimensions["B"].width = 48
    for row in summary.iter_rows(min_row=2, max_col=2):
        row[0].font = Font(bold=True, color="4A4A4F")
        row[1].alignment = Alignment(wrap_text=True, vertical="top")

    details = workbook.create_sheet("结果明细")
    details.sheet_view.showGridLines = False
    raw_columns = schema.get("columns", [])
    columns = [column for column in raw_columns if isinstance(column, dict)]
    headers = ["序号", *[str(column.get("label", column.get("key", "字段"))) for column in columns], "状态", "原因", "完成时间"]
    details.append(headers)
    for item in items:
        values = _load_values(item.values_json)
        row = [
            item.sequence,
            *[_safe_cell(values.get(str(column.get("key", "")))) for column in columns],
            ITEM_STATUS_LABELS.get(item.status, item.status),
            _safe_cell(item.reason),
            _format_timestamp_ms(item.reported_at_ms, timezone),
        ]
        details.append(row)
    _style_header(details, len(headers))
    details.freeze_panes = "A2"
    details.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, details.max_row)}"
    for index, header in enumerate(headers, start=1):
        width = 12
        if header in {"原因"}:
            width = 42
        elif header in {"完成时间"}:
            width = 22
        elif header not in {"序号", "状态"}:
            width = 20
        details.column_dimensions[get_column_letter(index)].width = width
    for row in details.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        status_cell = row[-3]
        fill = {
            "成功": "E6F4EA",
            "失败": "FDE8E7",
            "无数据": "EEF1F5",
            "已跳过": "FFF4DA",
            "未完成": "FDE8E7",
        }.get(str(status_cell.value), "FFFFFF")
        status_cell.fill = PatternFill("solid", fgColor=fill)

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _style_header(sheet: Any, column_count: int) -> None:
    fill = PatternFill("solid", fgColor="E9F2FD")
    font = Font(bold=True, color="174A7E")
    for cell in sheet[1][:column_count]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 24


def _format_timestamp(value: int | None, timezone: ZoneInfo) -> str:
    if value is None:
        return "—"
    return datetime.fromtimestamp(value, timezone).strftime("%Y-%m-%d %H:%M:%S %Z")


def _format_timestamp_ms(value: int, timezone: ZoneInfo) -> str:
    return datetime.fromtimestamp(value / 1000, timezone).strftime("%Y-%m-%d %H:%M:%S")


def _load_values(payload: str) -> dict[str, Any]:
    try:
        value = json.loads(payload)
    except (TypeError, ValueError):
        return {}
    return value if isinstance(value, dict) else {}


def _safe_cell(value: object) -> object:
    if value is None:
        return "—"
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value
