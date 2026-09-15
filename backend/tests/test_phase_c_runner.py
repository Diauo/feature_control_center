from __future__ import annotations

import io
import json
import time
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import undefer

from app.infrastructure.models import (
    AuditLogModel,
    RunEventModel,
    RunModel,
    RunReportItemModel,
    RunReportModel,
    SystemSettingModel,
)
from app.runner.dependencies import DependencyPreparer
from app.runner.materializer import RunMaterializer, ScriptIdentity
from app.runner.worker import RunnerWorker
from conftest import initialize
from test_phase_b_features import PASSWORD, upload


def executable_zip(
    name: str,
    source: str,
    *,
    default_source: bytes = b"queued-source",
    report: dict[str, object] | None = None,
) -> bytes:
    metadata = {
        "name": name,
        "description": "Runner 集成测试",
        "entrypoint": "run",
        "configs": {
            "api_token": {"type": "secret", "label": "接口密钥", "required": True},
        },
        "data_source": {"filename": "dataSource.txt", "required": True, "extensions": [".txt"]},
    }
    if report is not None:
        metadata["report"] = report
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("__init__.py", f"__meta__ = {metadata!r}\n{source}\n")
        archive.writestr("dataSource.txt", default_source)
    return buffer.getvalue()


def make_worker(app, runner_id: str = "phase-c-test") -> RunnerWorker:
    services = app.extensions["fcc_services"]
    return RunnerWorker(
        database=services.database,
        settings=services.settings,
        clock=services.clock,
        materializer=RunMaterializer(
            services.database,
            services.paths,
            services.runs.snapshot_cipher,
            ScriptIdentity(),
        ),
        dependency_preparer=DependencyPreparer(
            services.database,
            services.settings,
            services.clock,
            app.extensions["test_environment_builder"],
        ),
        runner_id=runner_id,
    )


def setup_feature(client, tmp_path, package: bytes):
    session = initialize(client, tmp_path / "data", password=PASSWORD)
    csrf = session["csrfToken"]
    customer_id = client.get("/api/customers").get_json()["items"][0]["id"]
    response = upload(client, csrf, customer_id, package, "runnable.zip")
    assert response.status_code == 201
    feature = client.get(f"/api/customers/{customer_id}/features").get_json()["items"][0]
    configured = client.put(
        f"/api/customer-features/{feature['id']}/config",
        headers={"X-CSRF-Token": csrf},
        json={"values": {"api_token": "snapshot-secret"}, "clearSecrets": []},
    )
    assert configured.status_code == 200
    return csrf, customer_id, feature


def wait_terminal(worker: RunnerWorker, services, request_id: str, timeout: float = 10) -> RunModel:
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        worker.tick()
        with services.database.session() as db:
            run = db.get(RunModel, request_id)
            if run and run.status in {"SUCCEEDED", "FAILED", "STOPPED", "TIMED_OUT", "INTERRUPTED"}:
                return run
        time.sleep(0.05)
    raise AssertionError("任务没有在预期时间内结束")


def test_runner_freezes_inputs_streams_logs_and_generates_final_download(client, app, tmp_path):
    source = """
import sys
def run(ctx):
    source = ctx.get_data_source_path().read_text(encoding='utf-8')
    ctx.log.info('snapshot verified', sourceValue=source, tokenMatches=ctx.config.get_secret('api_token') == 'snapshot-secret')
    print('plain stdout')
    print('plain stderr', file=sys.stderr)
"""
    csrf, customer_id, feature = setup_feature(client, tmp_path, executable_zip("快照与日志", source))
    created = client.post(f"/api/customer-features/{feature['id']}/runs", headers={"X-CSRF-Token": csrf})
    assert created.status_code == 202
    request_id = created.get_json()["run"]["requestId"]

    replaced = client.put(
        f"/api/customer-features/{feature['id']}/data-source",
        headers={"X-CSRF-Token": csrf},
        data={"file": (io.BytesIO(b"replacement-after-queue"), "replacement.txt")},
        content_type="multipart/form-data",
    )
    assert replaced.status_code == 200
    client.put(
        f"/api/customer-features/{feature['id']}/config",
        headers={"X-CSRF-Token": csrf},
        json={"values": {"api_token": "changed-after-queue"}, "clearSecrets": []},
    )

    services = app.extensions["fcc_services"]
    worker = make_worker(app)
    run = wait_terminal(worker, services, request_id)
    assert run.status == "SUCCEEDED"
    assert run.final_event_sequence == run.last_event_sequence
    assert not (services.paths.temp_dir / "runs" / request_id).exists()
    with services.database.session() as db:
        stored = db.scalar(
            select(RunModel).options(undefer(RunModel.config_snapshot_ciphertext)).where(RunModel.request_id == request_id)
        )
        assert b"snapshot-secret" not in stored.config_snapshot_ciphertext
        events = db.scalars(
            select(RunEventModel).where(RunEventModel.run_request_id == request_id).order_by(RunEventModel.sequence)
        ).all()
        assert [item.sequence for item in events] == list(range(1, len(events) + 1))
        messages = [item.message for item in events]
        assert "plain stdout" in messages
        assert "plain stderr" in messages
        sdk_event = next(item for item in events if item.message == "snapshot verified")
        context = json.loads(sdk_event.context_json)
        assert context == {"sourceValue": "queued-source", "tokenMatches": True}

    event_page = client.get(f"/api/runs/{request_id}/events?after=0").get_json()
    assert event_page["terminal"] is True
    assert event_page["latestSequence"] == len(event_page["items"])
    text_log = client.get(f"/api/runs/{request_id}/log.log")
    jsonl_log = client.get(f"/api/runs/{request_id}/log.jsonl")
    assert text_log.status_code == 200 and b"plain stdout" in text_log.data
    assert jsonl_log.status_code == 200
    assert all(json.loads(line)["requestId"] == request_id for line in jsonl_log.data.decode().splitlines())
    stream = client.get(f"/api/runs/{request_id}/events/stream?after=0")
    assert stream.status_code == 200
    assert b"event: log" in stream.data and b"event: run" in stream.data
    worker.shutdown()


def test_runner_persists_localized_structured_report_and_exports_xlsx(client, app, tmp_path):
    report_schema = {
        "title": "库存同步结果",
        "item_label": "商品",
        "columns": [
            {"key": "sku", "label": "商品编码", "type": "string"},
            {"key": "storeCode", "label": "店铺编码", "type": "string"},
            {"key": "quantity", "label": "库存", "type": "integer"},
        ],
    }
    source = """
def run(ctx):
    ctx.report.start(expected_total=3)
    ctx.report.success(values={'sku': 'SKU-001', 'storeCode': 'CK001', 'quantity': 12})
    ctx.report.failed(values={'sku': 'SKU-002', 'storeCode': 'CK002', 'quantity': 0}, reason='接口超时', reason_code='NETWORK.TIMEOUT')
    ctx.report.no_data(values={'sku': '=2+2', 'storeCode': 'CK003', 'quantity': 0}, reason='上游没有库存记录')
"""
    package = executable_zip("结构化报表", source, report=report_schema)
    csrf, _customer_id, feature = setup_feature(client, tmp_path, package)
    request_id = client.post(
        f"/api/customer-features/{feature['id']}/runs", headers={"X-CSRF-Token": csrf}
    ).get_json()["run"]["requestId"]
    services = app.extensions["fcc_services"]
    worker = make_worker(app, "report-test")
    run = wait_terminal(worker, services, request_id)
    assert run.status == "SUCCEEDED"

    report_response = client.get(f"/api/runs/{request_id}/report")
    assert report_response.status_code == 200
    summary = report_response.get_json()["report"]
    assert summary == {
        "title": "库存同步结果",
        "itemLabel": "商品",
        "columns": report_schema["columns"],
        "status": "COMPLETE",
        "expectedTotal": 3,
        "reportedTotal": 3,
        "total": 3,
        "successCount": 1,
        "failedCount": 1,
        "noDataCount": 1,
        "skippedCount": 0,
        "unfinishedCount": 0,
        "unreportedCount": 0,
        "validationErrorCount": 0,
        "startedAt": summary["startedAt"],
        "completedAt": summary["completedAt"],
        "detailsPurgedAt": None,
    }
    failed = client.get(f"/api/runs/{request_id}/report/items?status=FAILED&page=1&pageSize=20")
    assert failed.status_code == 200
    assert failed.get_json()["pagination"]["total"] == 1
    assert failed.get_json()["items"] == [
        {
            "sequence": 2,
            "status": "FAILED",
            "reason": "接口超时",
            "values": {"sku": "SKU-002", "storeCode": "CK002", "quantity": 0},
            "reportedAtMs": failed.get_json()["items"][0]["reportedAtMs"],
        }
    ]
    assert "reasonCode" not in failed.get_json()["items"][0]

    workbook_response = client.get(f"/api/runs/{request_id}/report.xlsx")
    assert workbook_response.status_code == 200
    workbook = load_workbook(io.BytesIO(workbook_response.data), read_only=True, data_only=False)
    assert workbook.sheetnames == ["执行汇总", "结果明细"]
    details = list(workbook["结果明细"].iter_rows(values_only=True))
    assert details[0] == ("序号", "商品编码", "店铺编码", "库存", "状态", "原因", "完成时间")
    assert details[1][4] == "成功"
    assert details[2][4:6] == ("失败", "接口超时")
    assert details[3][1] == "'=2+2"
    assert details[3][4] == "无数据"
    workbook.close()
    with services.database.session() as db:
        audit = db.scalar(
            select(AuditLogModel)
            .where(
                AuditLogModel.action == "run.report.download",
                AuditLogModel.target_id == request_id,
            )
            .order_by(AuditLogModel.occurred_at.desc())
        )
        assert audit is not None
        audit_details = json.loads(audit.details_json)
        assert audit_details == {"format": "xlsx", "itemCount": 3, "complete": True}
    worker.shutdown()


def test_report_limit_defaults_to_60000_and_is_snapshotted_per_run(client, app, tmp_path):
    report_schema = {
        "title": "大批量同步结果",
        "item_label": "记录",
        "columns": [{"key": "recordId", "label": "记录编号", "type": "string"}],
    }
    source = """
def run(ctx):
    ctx.report.start(expected_total=50000)
"""
    package = executable_zip("大批量报表", source, report=report_schema)
    csrf, _customer_id, feature = setup_feature(client, tmp_path, package)

    settings = client.get("/api/admin/settings")
    assert settings.status_code == 200
    assert settings.get_json()["values"]["runner.report_max_items"] == 60_000

    accepted_id = client.post(
        f"/api/customer-features/{feature['id']}/runs",
        headers={"X-CSRF-Token": csrf},
    ).get_json()["run"]["requestId"]
    worker = make_worker(app, "large-report-default")
    accepted = wait_terminal(worker, app.extensions["fcc_services"], accepted_id)
    assert accepted.status == "SUCCEEDED"
    accepted_report = client.get(f"/api/runs/{accepted_id}/report").get_json()["report"]
    assert accepted_report["expectedTotal"] == 50_000
    assert accepted_report["validationErrorCount"] == 0
    worker.shutdown()

    updated = client.put(
        "/api/admin/settings",
        headers={"X-CSRF-Token": csrf},
        json={"values": {"runner.report_max_items": 40_000}},
    )
    assert updated.status_code == 200
    assert updated.get_json()["values"]["runner.report_max_items"] == 40_000
    too_large = client.put(
        "/api/admin/settings",
        headers={"X-CSRF-Token": csrf},
        json={"values": {"runner.report_max_items": 1_000_001}},
    )
    assert too_large.status_code == 400
    assert too_large.get_json()["error"]["code"] == "INVALID_SETTING_VALUE"

    rejected_id = client.post(
        f"/api/customer-features/{feature['id']}/runs",
        headers={"X-CSRF-Token": csrf},
    ).get_json()["run"]["requestId"]
    worker = make_worker(app, "large-report-configured")
    rejected = wait_terminal(worker, app.extensions["fcc_services"], rejected_id)
    assert rejected.status == "FAILED"
    with app.extensions["fcc_services"].database.session() as db:
        messages = db.scalars(
            select(RunEventModel)
            .where(RunEventModel.run_request_id == rejected_id)
            .order_by(RunEventModel.sequence)
        ).all()
        assert any("expected_total 必须是 0 到 40000 的整数" in item.message for item in messages)
    worker.shutdown()


def test_stop_request_reaches_child_and_seals_stopped(client, app, tmp_path):
    source = """
import time
def run(ctx):
    ctx.report.start(expected_total=2)
    ctx.report.success(values={'sku': 'SKU-001'})
    ctx.log.info('loop started')
    while True:
        ctx.raise_if_cancelled()
        time.sleep(0.02)
"""
    report_schema = {
        "title": "停止测试结果",
        "item_label": "商品",
        "columns": [{"key": "sku", "label": "商品编码", "type": "string"}],
    }
    csrf, _customer_id, feature = setup_feature(
        client,
        tmp_path,
        executable_zip("可停止任务", source, report=report_schema),
    )
    request_id = client.post(
        f"/api/customer-features/{feature['id']}/runs", headers={"X-CSRF-Token": csrf}
    ).get_json()["run"]["requestId"]
    services = app.extensions["fcc_services"]
    worker = make_worker(app, "stop-test")
    deadline = time.monotonic() + 5
    while time.monotonic() < deadline:
        worker.tick()
        with services.database.session() as db:
            run = db.get(RunModel, request_id)
            report = db.get(RunReportModel, request_id)
            if run.status == "RUNNING" and report is not None and report.reported_total == 1:
                break
        time.sleep(0.03)
    stopped = client.post(f"/api/runs/{request_id}/stop", headers={"X-CSRF-Token": csrf})
    assert stopped.status_code == 200
    assert stopped.get_json()["run"]["status"] == "STOPPING"
    run = wait_terminal(worker, services, request_id)
    assert run.status == "STOPPED"
    assert run.stop_reason == "USER_REQUEST"
    report = client.get(f"/api/runs/{request_id}/report").get_json()["report"]
    assert report["status"] == "INCOMPLETE"
    assert report["successCount"] == 1
    assert report["unfinishedCount"] == 1
    worker.shutdown()


def test_runner_startup_marks_orphaned_active_run_interrupted(client, app, tmp_path):
    source = "def run(ctx):\n    ctx.log.info('not reached')"
    csrf, _customer_id, feature = setup_feature(client, tmp_path, executable_zip("恢复任务", source))
    request_id = client.post(
        f"/api/customer-features/{feature['id']}/runs", headers={"X-CSRF-Token": csrf}
    ).get_json()["run"]["requestId"]
    services = app.extensions["fcc_services"]
    with services.database.session() as db:
        run = db.get(RunModel, request_id)
        run.status = "RUNNING"
        run.runner_id = "dead-runner"
        run.process_id = 98765
    worker = make_worker(app, "recovery-test")
    worker.start()
    recovered = client.get(f"/api/runs/{request_id}").get_json()["run"]
    assert recovered["status"] == "INTERRUPTED"
    assert recovered["finalSequence"] == recovered["latestSequence"]
    worker.shutdown()


def test_runtime_limit_uses_two_stage_termination_and_marks_timeout(client, app, tmp_path):
    source = """
import time
def run(ctx):
    ctx.log.info('ignoring cooperative cancellation for test')
    while True:
        time.sleep(0.05)
"""
    csrf, _customer_id, feature = setup_feature(client, tmp_path, executable_zip("超时任务", source))
    policy = client.patch(
        f"/api/admin/customer-features/{feature['id']}",
        headers={"X-CSRF-Token": csrf},
        json={"maxRuntimeSeconds": 1},
    )
    assert policy.status_code == 200
    services = app.extensions["fcc_services"]
    with services.database.session() as db:
        grace = db.get(SystemSettingModel, "runner.stop_grace_seconds")
        grace.value_json = "1"
    request_id = client.post(
        f"/api/customer-features/{feature['id']}/runs", headers={"X-CSRF-Token": csrf}
    ).get_json()["run"]["requestId"]
    worker = make_worker(app, "timeout-test")
    run = wait_terminal(worker, services, request_id, timeout=8)
    assert run.status == "TIMED_OUT"
    assert run.stop_reason == "TIMEOUT"
    worker.shutdown()


def test_original_legacy_module_contract_remains_executable(client, app, tmp_path):
    # The repository sample is copied into a ZIP exactly as a user would register it.
    source_dir = Path(__file__).parents[2] / "features" / "example_module"
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for path in source_dir.rglob("*.py"):
            archive.write(path, path.relative_to(source_dir).as_posix())
    session = initialize(client, tmp_path / "data", password=PASSWORD)
    csrf = session["csrfToken"]
    customer_id = client.get("/api/customers").get_json()["items"][0]["id"]
    registered = upload(client, csrf, customer_id, buffer.getvalue(), "legacy.zip")
    assert registered.status_code == 201
    feature = client.get(f"/api/customers/{customer_id}/features").get_json()["items"][0]
    request_id = client.post(
        f"/api/customer-features/{feature['id']}/runs", headers={"X-CSRF-Token": csrf}
    ).get_json()["run"]["requestId"]
    services = app.extensions["fcc_services"]
    worker = make_worker(app, "legacy-test")
    run = wait_terminal(worker, services, request_id)
    assert run.status == "SUCCEEDED"
    messages = [item["message"] for item in client.get(f"/api/runs/{request_id}/events").get_json()["items"]]
    assert "开始执行示例模块" in messages
    assert "执行成功" in messages
    worker.shutdown()


def test_legacy_context_log_shortcuts_keep_their_levels(client, app, tmp_path):
    source = """
def run(ctx):
    ctx.debug('legacy debug', sku='SKU-1')
    ctx.info('legacy info')
    ctx.warning('legacy warning')
    ctx.error('legacy error', stage='request')
"""
    csrf, _customer_id, feature = setup_feature(client, tmp_path, executable_zip("旧日志快捷方法", source))
    request_id = client.post(
        f"/api/customer-features/{feature['id']}/runs", headers={"X-CSRF-Token": csrf}
    ).get_json()["run"]["requestId"]
    services = app.extensions["fcc_services"]
    worker = make_worker(app, "legacy-level-test")
    run = wait_terminal(worker, services, request_id)
    assert run.status == "SUCCEEDED"
    events = client.get(f"/api/runs/{request_id}/events").get_json()["items"]
    levels = {item["message"]: item["level"] for item in events}
    assert levels["legacy debug"] == "DEBUG"
    assert levels["legacy info"] == "INFO"
    assert levels["legacy warning"] == "WARNING"
    assert levels["legacy error"] == "ERROR"
    worker.shutdown()


def test_runner_purges_expired_logs_but_keeps_run_summary(client, app, tmp_path):
    source = """
def run(ctx):
    ctx.log.info('retention payload', sku='SKU-001')
    ctx.report.start(expected_total=1)
    ctx.report.success(values={'sku': 'SKU-001'})
"""
    report_schema = {
        "title": "保留策略测试",
        "item_label": "商品",
        "columns": [{"key": "sku", "label": "商品编码", "type": "string"}],
    }
    csrf, _customer_id, feature = setup_feature(
        client,
        tmp_path,
        executable_zip("日志保留测试", source, report=report_schema),
    )
    request_id = client.post(
        f"/api/customer-features/{feature['id']}/runs", headers={"X-CSRF-Token": csrf}
    ).get_json()["run"]["requestId"]
    services = app.extensions["fcc_services"]
    worker = make_worker(app, "retention-run-test")
    completed = wait_terminal(worker, services, request_id)
    assert completed.status == "SUCCEEDED"
    worker.shutdown()

    with services.database.session() as db:
        run = db.get(RunModel, request_id)
        run.finished_at = services.clock.now() - 180 * 86_400
        event_count = len(
            db.scalars(select(RunEventModel).where(RunEventModel.run_request_id == request_id)).all()
        )
        assert event_count > 0

    disabled = client.put(
        "/api/admin/settings",
        headers={"X-CSRF-Token": csrf},
        json={"values": {"logs.retention_days": 0}},
    )
    assert disabled.status_code == 200
    cleaner = make_worker(app, "retention-disabled-test")
    assert cleaner.log_retention.purge_batch().purged_runs == 0
    with services.database.session() as db:
        assert db.get(RunModel, request_id).logs_purged_at is None
    cleaner.shutdown()

    enabled = client.put(
        "/api/admin/settings",
        headers={"X-CSRF-Token": csrf},
        json={"values": {"logs.retention_days": 180}},
    )
    assert enabled.status_code == 200
    assert enabled.get_json()["values"]["logs.retention_days"] == 180
    cleaner = make_worker(app, "retention-enabled-test")
    cleaner.tick()

    with services.database.session() as db:
        run = db.get(RunModel, request_id)
        assert run is not None
        assert run.status == "SUCCEEDED"
        assert run.logs_purged_at == services.clock.now()
        assert run.last_event_sequence == run.final_event_sequence
        assert db.scalars(
            select(RunEventModel).where(RunEventModel.run_request_id == request_id)
        ).first() is None
        report = db.get(RunReportModel, request_id)
        assert report is not None
        assert report.status == "COMPLETE"
        assert report.details_purged_at == services.clock.now()
        assert report.success_count == 1
        assert db.scalars(
            select(RunReportItemModel).where(RunReportItemModel.run_request_id == request_id)
        ).first() is None

    detail = client.get(f"/api/runs/{request_id}")
    assert detail.status_code == 200
    assert detail.get_json()["run"]["logsPurgedAt"] == services.clock.now()
    event_page = client.get(f"/api/runs/{request_id}/events?after=0")
    assert event_page.status_code == 200
    assert event_page.get_json()["items"] == []
    assert event_page.get_json()["logsPurgedAt"] == services.clock.now()
    for extension in ("log", "jsonl"):
        response = client.get(f"/api/runs/{request_id}/log.{extension}")
        assert response.status_code == 410
        assert response.get_json()["error"]["code"] == "RUN_LOGS_PURGED"
    report = client.get(f"/api/runs/{request_id}/report")
    assert report.status_code == 200
    assert report.get_json()["report"]["successCount"] == 1
    assert report.get_json()["report"]["detailsPurgedAt"] == services.clock.now()
    report_items = client.get(f"/api/runs/{request_id}/report/items")
    assert report_items.status_code == 410
    assert report_items.get_json()["error"]["code"] == "RUN_REPORT_DETAILS_PURGED"
    assert client.get(f"/api/runs/{request_id}/report.xlsx").status_code == 410
    cleaner.shutdown()
